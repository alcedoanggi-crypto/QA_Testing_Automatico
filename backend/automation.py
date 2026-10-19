"""Pruebas automáticas del sistema y programador periódico.

La aplicación se prueba a sí misma: cada caso ejecuta peticiones reales contra la API
(mediante el cliente de pruebas de Flask, sin abrir puertos), valida la respuesta y el
resultado queda registrado en la base de datos como una ejecución por suite.
"""
import datetime as dt
import io
import os
import threading
import time
import uuid

import jwt
from openpyxl import load_workbook

from auth import ALGORITHM, SECRET_KEY
from logger import log_scheduler_error, log_test_batch_finished, log_test_batch_started, log_test_case_failed
from reports_engine import TRIGGER_SCHEDULED, build_run_record, repository

RUNNER_EMAIL = os.environ.get("QA_RUNNER_EMAIL", "admin@test.com")
RUNNER_PASSWORD = os.environ.get("QA_RUNNER_PASSWORD", "SecurePassword123")
SYSTEM_USER = "Sistema automático"

LOGIN_URL = "/api/v1/auth/login"
METRICS_URL = "/api/v1/dashboard/metrics"
EXPORT_URL = "/api/v1/reports/export"

_run_lock = threading.Lock()


class RunnerBusyError(RuntimeError):
    """Ya hay un lote de pruebas en ejecución."""


def is_running() -> bool:
    return _run_lock.locked()


def expect(condition, message):
    if not condition:
        raise AssertionError(message)


def expect_status(response, expected):
    expect(response.status_code == expected,
           f"Se esperaba HTTP {expected} y se obtuvo HTTP {response.status_code}")


class _Context:
    """Estado compartido entre casos de un mismo lote (cliente HTTP y token)."""

    def __init__(self, client):
        self.client = client
        self.token = None

    def auth(self, token=None):
        token = token or self.token
        expect(token, "No hay token disponible: el caso de login exitoso falló")
        return {"Authorization": f"Bearer {token}"}


# ----------------------------------------------------------------------------- Autenticación
def case_login_success(ctx):
    response = ctx.client.post(LOGIN_URL, json={"email": RUNNER_EMAIL, "password": RUNNER_PASSWORD})
    expect_status(response, 200)
    body = response.get_json()
    expect(len(str(body.get("token", "")).split(".")) == 3, "El token devuelto no tiene formato JWT")
    expect({"name", "role"} <= set(body.get("user") or {}), "La respuesta no incluye nombre y rol del usuario")
    ctx.token = body["token"]


def case_login_wrong_password(ctx):
    response = ctx.client.post(LOGIN_URL, json={"email": RUNNER_EMAIL, "password": "clave-incorrecta"})
    expect_status(response, 401)
    expect(response.get_json().get("error") == "Credenciales inválidas", "Mensaje de error inesperado")


def case_login_unknown_user(ctx):
    response = ctx.client.post(LOGIN_URL, json={"email": "no-existe@test.com", "password": "x"})
    expect_status(response, 401)


def case_login_missing_fields(ctx):
    response = ctx.client.post(LOGIN_URL, json={"email": ""})
    expect_status(response, 400)


def case_login_non_json_body(ctx):
    response = ctx.client.post(LOGIN_URL, data="email=admin", content_type="text/plain")
    expect_status(response, 400)


# ----------------------------------------------------------------------------- Seguridad
def case_metrics_without_token(ctx):
    expect_status(ctx.client.get(METRICS_URL), 401)


def case_token_wrong_signature(ctx):
    forged = jwt.encode({"sub": RUNNER_EMAIL, "exp": dt.datetime.now(dt.timezone.utc) + dt.timedelta(minutes=5)},
                        "clave-de-un-atacante-con-longitud-suficiente", algorithm=ALGORITHM)
    expect_status(ctx.client.get(METRICS_URL, headers=ctx.auth(forged)), 401)


def case_token_expired(ctx):
    expired = jwt.encode({"sub": RUNNER_EMAIL, "exp": dt.datetime.now(dt.timezone.utc) - dt.timedelta(minutes=1)},
                         SECRET_KEY, algorithm=ALGORITHM)
    response = ctx.client.get(METRICS_URL, headers=ctx.auth(expired))
    expect_status(response, 401)
    expect(response.get_json().get("error") == "El token ha expirado", "No se informó la expiración del token")


def case_wrong_auth_scheme(ctx):
    response = ctx.client.get(METRICS_URL, headers={"Authorization": f"Basic {ctx.token or 'x'}"})
    expect_status(response, 401)


def case_export_requires_token(ctx):
    expect_status(ctx.client.get(f"{EXPORT_URL}?format=csv"), 401)


# ----------------------------------------------------------------------------- Dashboard
def case_metrics_ok(ctx):
    response = ctx.client.get(METRICS_URL, headers=ctx.auth())
    expect_status(response, 200)
    body = response.get_json()
    missing = {"total_tests", "pass_rate", "bugs_found", "avg_response_time_ms", "recent_executions"} - set(body)
    expect(not missing, f"Faltan campos en la respuesta: {', '.join(sorted(missing))}")


def case_metrics_pass_rate_range(ctx):
    body = ctx.client.get(METRICS_URL, headers=ctx.auth()).get_json()
    expect(0 <= body["pass_rate"] <= 100, f"Pass rate fuera de rango: {body['pass_rate']}")


def case_metrics_consistency(ctx):
    body = ctx.client.get(METRICS_URL, headers=ctx.auth()).get_json()
    expect(body["passed"] + body["failed"] + body["skipped"] == body["total_tests"],
           "Exitosas + Fallidas + Omitidas no coincide con el total de pruebas")
    expect(body["total_executions"] == len(body["recent_executions"]),
           "El total de ejecuciones no coincide con el listado")


# ----------------------------------------------------------------------------- Reportes
def case_export_csv(ctx):
    response = ctx.client.get(f"{EXPORT_URL}?format=csv", headers=ctx.auth())
    expect_status(response, 200)
    header = response.data.decode("utf-8-sig").splitlines()[0]
    expect("ID Ejecución" in header and "Pass Rate (%)" in header, "Cabecera del CSV incorrecta")


def case_export_excel(ctx):
    response = ctx.client.get(f"{EXPORT_URL}?format=excel", headers=ctx.auth())
    expect_status(response, 200)
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    expect({"Resumen", "Ejecuciones"} <= set(workbook.sheetnames), "El Excel no contiene las hojas esperadas")
    workbook.close()


def case_export_pdf(ctx):
    response = ctx.client.get(f"{EXPORT_URL}?format=pdf", headers=ctx.auth())
    expect_status(response, 200)
    expect(response.data[:4] == b"%PDF", "El archivo no tiene la firma de un PDF")
    expect("reporte_pruebas.pdf" in response.headers.get("Content-Disposition", ""), "Nombre de archivo incorrecto")


def case_export_invalid_format(ctx):
    expect_status(ctx.client.get(f"{EXPORT_URL}?format=docx", headers=ctx.auth()), 400)


# ----------------------------------------------------------------------------- API REST y vistas
def case_health(ctx):
    response = ctx.client.get("/api/v1/health")
    expect_status(response, 200)
    expect(response.get_json() == {"status": "ok"}, "Respuesta de health inesperada")


def case_login_view(ctx):
    response = ctx.client.get("/login")
    expect_status(response, 200)
    for element_id in ("email", "password", "btn-login", "login-error-msg"):
        expect(f'id="{element_id}"'.encode() in response.data, f"La vista de login no contiene #{element_id}")


def case_dashboard_view(ctx):
    response = ctx.client.get("/dashboard")
    expect_status(response, 200)
    for element_id in ("btn-export-pdf", "btn-export-excel", "btn-export-csv", "btn-logout", "table-test-runs"):
        expect(f'id="{element_id}"'.encode() in response.data, f"El dashboard no contiene #{element_id}")


def case_unknown_endpoint(ctx):
    response = ctx.client.get("/api/v1/no-existe")
    expect_status(response, 404)
    expect(response.is_json, "El 404 de la API no responde en JSON")


SUITES = [
    ("Autenticación", [
        ("Login con credenciales válidas devuelve 200 y un JWT", case_login_success),
        ("Login con contraseña incorrecta devuelve 401", case_login_wrong_password),
        ("Login con usuario inexistente devuelve 401", case_login_unknown_user),
        ("Login sin campos obligatorios devuelve 400", case_login_missing_fields),
        ("Login con cuerpo no JSON devuelve 400", case_login_non_json_body),
    ]),
    ("Seguridad", [
        ("Métricas sin token devuelven 401", case_metrics_without_token),
        ("Token con firma inválida es rechazado", case_token_wrong_signature),
        ("Token expirado es rechazado", case_token_expired),
        ("Esquema de autorización distinto a Bearer es rechazado", case_wrong_auth_scheme),
        ("Exportación sin token devuelve 401", case_export_requires_token),
    ]),
    ("Dashboard", [
        ("Métricas con token devuelven 200 y todos los campos", case_metrics_ok),
        ("Pass rate está entre 0 y 100", case_metrics_pass_rate_range),
        ("Totales de métricas son consistentes", case_metrics_consistency),
    ]),
    ("Reportes", [
        ("Exportación CSV con cabecera correcta", case_export_csv),
        ("Exportación Excel con hojas Resumen y Ejecuciones", case_export_excel),
        ("Exportación PDF válida", case_export_pdf),
        ("Formato de exportación inválido devuelve 400", case_export_invalid_format),
    ]),
    ("API REST", [
        ("Health check responde 200", case_health),
        ("Vista de login contiene los elementos requeridos", case_login_view),
        ("Vista de dashboard contiene los elementos requeridos", case_dashboard_view),
        ("Endpoint inexistente devuelve 404 en JSON", case_unknown_endpoint),
    ]),
]


def run_system_tests(app, executed_by=SYSTEM_USER, trigger=TRIGGER_SCHEDULED, environment="Local") -> dict:
    """Ejecuta todas las suites contra la aplicación y registra una ejecución por suite.

    Lanza ``RunnerBusyError`` si ya hay un lote en curso.
    """
    if not _run_lock.acquire(blocking=False):
        raise RunnerBusyError("Ya hay una ejecución de pruebas en curso")
    try:
        batch_id = f"LOTE-{dt.datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:4].upper()}"
        log_test_batch_started(batch_id, trigger, executed_by)
        ctx = _Context(app.test_client())

        suite_results = []
        for suite, cases in SUITES:
            results = []
            for name, func in cases:
                started = time.perf_counter()
                try:
                    func(ctx)
                    status, detail = "Passed", ""
                except AssertionError as exc:
                    status, detail = "Failed", str(exc)
                except Exception as exc:  # un error inesperado también es un fallo del caso
                    status, detail = "Failed", f"Error inesperado: {exc!r}"
                elapsed_ms = round((time.perf_counter() - started) * 1000, 2)
                if status == "Failed":
                    log_test_case_failed(batch_id, suite, name, detail)
                results.append({"name": name, "status": status, "response_ms": elapsed_ms, "detail": detail})
            suite_results.append((suite, results))

        # Se registra al final para que los casos de métricas vean un estado consistente.
        runs = []
        for suite, results in suite_results:
            times = [r["response_ms"] for r in results]
            record = build_run_record(
                suite=suite, environment=environment, total=len(results),
                failed=sum(1 for r in results if r["status"] == "Failed"), skipped=0,
                avg_response_ms=sum(times) / len(times), duration_s=sum(times) / 1000,
                executed_by=executed_by, trigger=trigger, batch_id=batch_id,
            )
            runs.append(repository.add_run(record, results))

        total = sum(r["total"] for r in runs)
        failed = sum(r["failed"] for r in runs)
        log_test_batch_finished(batch_id, total, failed)
        return {"batch_id": batch_id, "total": total, "passed": total - failed, "failed": failed, "runs": runs}
    finally:
        _run_lock.release()


class AutoTestScheduler:
    """Hilo en segundo plano que ejecuta las pruebas al iniciar y luego cada ``interval_minutes``."""

    def __init__(self, app, interval_minutes: float, initial_delay_seconds: float = 3):
        self.app = app
        self.interval_minutes = interval_minutes
        self.initial_delay_seconds = initial_delay_seconds
        self.next_run_at = None
        self.last_error = None
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._loop, name="qa-auto-tests", daemon=True)

    def start(self):
        self._schedule(self.initial_delay_seconds)
        self._thread.start()

    def stop(self):
        self._stop.set()
        self._thread.join(timeout=10)

    def _schedule(self, seconds):
        self.next_run_at = dt.datetime.now() + dt.timedelta(seconds=seconds)

    def _loop(self):
        delay = self.initial_delay_seconds
        while not self._stop.wait(delay):
            try:
                run_system_tests(self.app)
                self.last_error = None
            except RunnerBusyError:
                pass  # alguien lanzó un lote manual justo ahora; se espera al siguiente ciclo
            except Exception as exc:
                self.last_error = str(exc)
                log_scheduler_error(repr(exc))
            delay = self.interval_minutes * 60
            self._schedule(delay)

    def status(self) -> dict:
        return {
            "enabled": True,
            "interval_minutes": self.interval_minutes,
            "next_run_at": self.next_run_at.strftime("%Y-%m-%d %H:%M:%S") if self.next_run_at else None,
            "running": is_running(),
            "last_error": self.last_error,
        }


def automation_status(scheduler) -> dict:
    if scheduler is None:
        return {"enabled": False, "interval_minutes": 0, "next_run_at": None, "running": is_running(),
                "last_error": None}
    return scheduler.status()
