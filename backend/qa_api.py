"""API REST de los módulos de QA Center: análisis de URL, resultados externos, novedades, programadores,
aplicaciones y configuración. Todas las rutas exigen JWT."""
import csv
import io
import os
import re

from flask import Blueprint, current_app, g, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import monitor
from auth import token_required
from importers import parse_report
from scan_report import build_scan_report
from logger import get_logger
from notifier import NotificationError, available_channels, build_message, build_subject, send_email, send_webhook
from qa_store import (
    INCIDENT_SOURCES, INCIDENT_STATUSES, PRIORITIES, SEVERITIES, NotFoundError, clean_url, store,
)
from reports_engine import (
    ENVIRONMENTS, TOOLS, TRIGGER_IMPORTED, build_run_record, parse_datetime, repository, update_manual_run,
)

api = Blueprint("qa_api", __name__, url_prefix="/api/v1")


def current_user_name() -> str:
    return g.current_user.get("name", g.current_user["sub"])


def body() -> dict:
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ValueError("El cuerpo de la petición debe ser JSON")
    return data


@api.errorhandler(ValueError)
def handle_value_error(exc):
    return jsonify({"error": str(exc)}), 400


@api.errorhandler(NotFoundError)
def handle_not_found(exc):
    return jsonify({"error": str(exc)}), 404


# ----------------------------------------------------------------------------- catálogos
@api.get("/catalogs")
@token_required
def catalogs():
    return jsonify({
        "environments": ENVIRONMENTS, "tools": TOOLS, "severities": SEVERITIES, "priorities": PRIORITIES,
        "incident_statuses": INCIDENT_STATUSES, "incident_sources": INCIDENT_SOURCES,
    })


# ----------------------------------------------------------------------------- programadores
@api.get("/developers")
@token_required
def list_developers():
    return jsonify(store.list_developers())


@api.post("/developers")
@token_required
def create_developer():
    return jsonify(store.create_developer(body())), 201


@api.put("/developers/<int:dev_id>")
@token_required
def update_developer(dev_id):
    return jsonify(store.update_developer(dev_id, body()))


@api.delete("/developers/<int:dev_id>")
@token_required
def delete_developer(dev_id):
    store.delete_developer(dev_id)
    return jsonify({"deleted": dev_id})


# ----------------------------------------------------------------------------- aplicaciones
@api.get("/targets")
@token_required
def list_targets():
    return jsonify(store.list_targets())


@api.post("/targets")
@token_required
def create_target():
    return jsonify(store.create_target(body())), 201


@api.put("/targets/<int:target_id>")
@token_required
def update_target(target_id):
    return jsonify(store.update_target(target_id, body()))


@api.delete("/targets/<int:target_id>")
@token_required
def delete_target(target_id):
    store.delete_target(target_id)
    return jsonify({"deleted": target_id})


@api.get("/monitor")
@token_required
def monitor_status():
    scheduler = current_app.extensions.get("qa_monitor")
    if scheduler is None:
        return jsonify({"enabled": False, "targets": len(store.monitored_targets())})
    return jsonify(scheduler.status())


@api.post("/monitor/run")
@token_required
def monitor_run_now():
    scheduler = current_app.extensions.get("qa_monitor")
    targets = store.monitored_targets()
    if not targets:
        raise ValueError("No hay aplicaciones marcadas para monitorear")
    if scheduler is not None:
        scheduler.run_now()
        return jsonify({"message": f"Monitoreo de {len(targets)} aplicación(es) en cola"}), 202
    scans = [monitor.start_scan(t["url"], current_user_name(), target_id=t["id"], max_pages=t["max_pages"],
                                trigger=monitor.TRIGGER_MONITOR) for t in targets]
    return jsonify({"message": f"{len(scans)} análisis iniciados", "scans": scans}), 202


# ----------------------------------------------------------------------------- análisis de URL
@api.post("/scans")
@token_required
def create_scan():
    data = body()
    target_id = data.get("target_id") or None
    if target_id:
        target = store.get_target(int(target_id))
        url = clean_url(data.get("url") or target["url"])
        default_pages = target["max_pages"]
    else:
        url = clean_url(data.get("url"))
        default_pages = 1
    try:
        max_pages = int(data.get("max_pages") or default_pages)
        wait_seconds = float(data["wait_seconds"]) if data.get("wait_seconds") not in (None, "") else None
    except (TypeError, ValueError):
        raise ValueError("Páginas y espera deben ser numéricos") from None
    if not 1 <= max_pages <= 20:
        raise ValueError("Páginas a recorrer debe estar entre 1 y 20")
    if wait_seconds is not None and not 0 <= wait_seconds <= 30:
        raise ValueError("La espera debe estar entre 0 y 30 segundos")
    auto = data.get("auto_incidents")
    scan = monitor.start_scan(url, current_user_name(), target_id=int(target_id) if target_id else None,
                              max_pages=max_pages, wait_seconds=wait_seconds,
                              auto_incidents=None if auto is None else bool(auto))
    return jsonify(scan), 202


@api.get("/scans")
@token_required
def list_scans():
    return jsonify(store.list_scans())


@api.get("/scans/<scan_id>")
@token_required
def get_scan(scan_id):
    return jsonify(store.get_scan(scan_id, with_issues=True))


@api.delete("/scans/<scan_id>")
@token_required
def delete_scan(scan_id):
    if store.get_scan(scan_id)["status"] == "En curso":
        raise ValueError("No se puede eliminar un análisis en curso")
    store.delete_scan(scan_id)
    return jsonify({"deleted": scan_id})


@api.get("/scans/<scan_id>/screenshot")
@token_required
def scan_screenshot(scan_id):
    store.get_scan(scan_id)
    path = os.path.join(store.screenshots_dir, f"{scan_id}.png")
    if not os.path.isfile(path):
        return jsonify({"error": "El análisis no tiene captura"}), 404
    return send_file(path, mimetype="image/png", max_age=0)


@api.get("/scans/<scan_id>/report")
@token_required
def scan_report(scan_id):
    scan = store.get_scan(scan_id, with_issues=True)
    if scan["status"] == "En curso":
        raise ValueError("El análisis aún está en curso")
    shot = os.path.join(store.screenshots_dir, f"{scan_id}.png")
    buffer, mimetype, filename = build_scan_report(scan, request.args.get("format", "pdf"), shot)
    return send_file(buffer, mimetype=mimetype, as_attachment=True, download_name=filename, max_age=0)


@api.get("/reports/auto")
@token_required
def auto_reports():
    settings = store.get_settings()
    return jsonify({"reports": monitor.list_auto_reports(), "report_hour": settings["report_hour"],
                    "report_recipients": settings["report_recipients"],
                    "last_date": settings["report_last_date"]})


@api.post("/reports/auto")
@token_required
def generate_auto_report():
    return jsonify(monitor.generate_auto_report(current_user_name())), 201


@api.get("/reports/auto/<name>")
@token_required
def download_auto_report(name):
    if not re.fullmatch(r"reporte_qa_[\d_-]+\.(pdf|xlsx)", name):
        raise NotFoundError("Reporte no encontrado")
    path = os.path.join(store.reports_dir, name)
    if not os.path.isfile(path):
        raise NotFoundError("Reporte no encontrado")
    return send_file(path, as_attachment=True, download_name=name, max_age=0)


@api.delete("/reports/auto/<name>")
@token_required
def delete_auto_report(name):
    if not re.fullmatch(r"reporte_qa_[\d_-]+\.(pdf|xlsx)", name):
        raise NotFoundError("Reporte no encontrado")
    path = os.path.join(store.reports_dir, name)
    if not os.path.isfile(path):
        raise NotFoundError("Reporte no encontrado")
    os.remove(path)
    return jsonify({"deleted": name})


@api.post("/scans/<scan_id>/incidents")
@token_required
def scan_to_incidents(scan_id):
    data = body()
    ids = data.get("issue_ids") or []
    if not ids:
        raise ValueError("Selecciona al menos un error")
    summary = monitor.create_incidents_from_issues(
        scan_id, ids, current_user_name(), developer_id=data.get("developer_id") or None,
        notify=bool(data.get("notify")))
    return jsonify(summary), 201


# ----------------------------------------------------------------------------- resultados de pruebas
@api.get("/test-runs")
@token_required
def list_runs():
    runs = repository.get_runs()
    origin = request.args.get("origin")
    if origin == "external":
        runs = [r for r in runs if r["batch_id"] is None]
    return jsonify(runs)


@api.get("/test-runs/<run_id>")
@token_required
def get_run(run_id):
    run = repository.get_run(run_id)
    if run is None:
        raise NotFoundError("Ejecución no encontrada")
    return jsonify({**run, "cases": repository.get_run_cases(run_id)})


@api.put("/test-runs/<run_id>")
@token_required
def update_run(run_id):
    run = update_manual_run(run_id, body())
    if run is None:
        raise NotFoundError("Ejecución no encontrada")
    get_logger().info("EJECUCION_EDITADA - ID: %s | Usuario: %s", run_id, g.current_user["sub"])
    return jsonify(run)


@api.delete("/test-runs/<run_id>")
@token_required
def delete_run(run_id):
    if not repository.delete_run(run_id):
        raise NotFoundError("Ejecución no encontrada")
    get_logger().info("EJECUCION_ELIMINADA - ID: %s | Usuario: %s", run_id, g.current_user["sub"])
    return jsonify({"deleted": run_id})


@api.post("/test-runs/<run_id>/incidents")
@token_required
def run_to_incidents(run_id):
    run = repository.get_run(run_id)
    if run is None:
        raise NotFoundError("Ejecución no encontrada")
    data = request.get_json(silent=True) or {}
    developer_id = data.get("developer_id") or None
    cases = repository.get_run_cases(run_id)
    if cases:
        created = monitor.create_incidents_from_run(run, cases, current_user_name(), developer_id)
    else:
        if run["failed"] == 0:
            raise ValueError("La ejecución no tiene pruebas fallidas")
        created = [store.create_incident({
            "title": f"[{run['tool'] or 'Pruebas'}] {run['failed']} prueba(s) fallida(s) en {run['suite']}"[:200],
            "description": f"Ejecución {run_id}: {run['passed']}/{run['total']} exitosas, {run['failed']} fallidas.",
            "actual": run.get("notes") or "",
            "severity": "Media", "environment": run["environment"], "module": run.get("module") or run["suite"],
            "url": run.get("url") or "", "developer_id": developer_id,
        }, current_user_name(), source="Resultado de pruebas", source_ref=run_id)]
    return jsonify({"created": created}), 201


@api.post("/test-runs/import")
@token_required
def import_run():
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        raise ValueError("Adjunta el archivo de resultados (JSON de Newman o XML JUnit)")
    parsed = parse_report(upload.filename, upload.read())
    form = request.form
    tool = parsed["tool"] or form.get("tool") or "Selenium"
    if tool not in TOOLS:
        raise ValueError(f"Herramienta inválida. Use: {', '.join(TOOLS)}")
    environment = form.get("environment") or "QA"
    if environment not in ENVIRONMENTS:
        raise ValueError(f"Ambiente inválido. Use: {', '.join(ENVIRONMENTS)}")
    cases = parsed["cases"]
    failed = sum(1 for c in cases if c["status"] == "Failed")
    skipped = sum(1 for c in cases if c["status"] == "Skipped")
    url = clean_url(form.get("url"), required=False)
    record = build_run_record(
        suite=(form.get("suite") or parsed["suite"]).strip()[:40], environment=environment, total=len(cases),
        failed=failed, skipped=skipped, avg_response_ms=parsed["avg_response_ms"], duration_s=parsed["duration_s"],
        executed_by=current_user_name(), trigger=TRIGGER_IMPORTED, tool=tool,
        module=(form.get("module") or "").strip()[:80], url=url,
        notes=(form.get("notes") or f"Importado desde {upload.filename}").strip()[:4000],
        executed_at=parse_datetime(form.get("executed_at")),
    )
    run = repository.add_run(record, cases)
    created = []
    if form.get("create_incidents") in ("1", "true", "on") and failed:
        created = monitor.create_incidents_from_run(run, cases, current_user_name(),
                                                    form.get("developer_id") or None)
    get_logger().info("RESULTADOS_IMPORTADOS - ID: %s | Herramienta: %s | Casos: %d | Fallidos: %d | Usuario: %s",
                      run["run_id"], tool, len(cases), failed, g.current_user["sub"])
    return jsonify({**run, "cases": len(cases), "incidents_created": [i["code"] for i in created]}), 201


# ----------------------------------------------------------------------------- novedades
@api.get("/incidents")
@token_required
def list_incidents():
    return jsonify(store.list_incidents(request.args.to_dict()))


@api.post("/incidents")
@token_required
def create_incident():
    data = body()
    incident = store.create_incident(data, current_user_name())
    if data.get("notify") and incident.get("developer_id"):
        channels = available_channels(store.get_settings(include_secrets=True))
        if channels:
            monitor.send_incident(incident["id"], channels, "", current_user_name())
    get_logger().info("NOVEDAD_REGISTRADA - %s | %s | Severidad: %s", incident["code"], incident["title"],
                      incident["severity"])
    return jsonify(store.get_incident(incident["id"], with_history=True)), 201


@api.get("/incidents/<int:incident_id>")
@token_required
def get_incident(incident_id):
    incident = store.get_incident(incident_id, with_history=True)
    settings = store.get_settings()
    incident["share"] = {
        "subject": build_subject(incident),
        "message": build_message(incident, settings),
        "channels": available_channels(store.get_settings(include_secrets=True)),
    }
    return jsonify(incident)


@api.put("/incidents/<int:incident_id>")
@token_required
def update_incident(incident_id):
    return jsonify(store.update_incident(incident_id, body(), current_user_name()))


@api.delete("/incidents/<int:incident_id>")
@token_required
def delete_incident(incident_id):
    store.delete_incident(incident_id)
    get_logger().info("NOVEDAD_ELIMINADA - ID: %s | Usuario: %s", incident_id, g.current_user["sub"])
    return jsonify({"deleted": incident_id})


@api.post("/incidents/<int:incident_id>/comments")
@token_required
def comment_incident(incident_id):
    store.get_incident(incident_id)
    text = str(body().get("text") or "").strip()
    if not text:
        raise ValueError("Escribe el comentario")
    if len(text) > 4000:
        raise ValueError("El comentario admite máximo 4000 caracteres")
    return jsonify(store.add_history(incident_id, "Comentario", text, current_user_name())), 201


@api.post("/incidents/<int:incident_id>/notify")
@token_required
def notify(incident_id):
    data = body()
    channels = [c for c in data.get("channels") or [] if c in ("email", "webhook")]
    note = str(data.get("note") or "").strip()[:2000]
    manual = data.get("manual_channel")
    if manual:
        return jsonify({"history": monitor.record_manual_share(incident_id, str(manual)[:40],
                                                               current_user_name())})
    if not channels:
        raise ValueError("Elige al menos un canal de envío")
    results = monitor.send_incident(incident_id, channels, note, current_user_name())
    ok = any(r == "ok" for r in results.values())
    return jsonify({"results": results}), 200 if ok else 502


@api.post("/incidents/bulk")
@token_required
def bulk_incidents():
    """Acciones masivas: asignar, cambiar estado, eliminar o enviar varias novedades a la vez."""
    data = body()
    ids = [int(i) for i in data.get("ids") or []]
    action = data.get("action")
    if not ids:
        raise ValueError("Selecciona al menos una novedad")
    author = current_user_name()
    done = 0
    for incident_id in ids:
        if action == "delete":
            store.delete_incident(incident_id)
        elif action == "assign":
            store.update_incident(incident_id, {"developer_id": data.get("developer_id") or None}, author)
        elif action == "status":
            store.update_incident(incident_id, {"status": data.get("status")}, author)
        elif action == "notify":
            channels = available_channels(store.get_settings(include_secrets=True))
            if not channels:
                raise ValueError("Configura un correo SMTP o un webhook para enviar automáticamente")
            monitor.send_incident(incident_id, channels, str(data.get("note") or ""), author)
        else:
            raise ValueError("Acción no soportada")
        done += 1
    return jsonify({"updated": done})


INCIDENT_EXPORT_COLUMNS = [
    ("code", "Código"), ("title", "Título"), ("severity", "Severidad"), ("priority", "Prioridad"),
    ("status", "Estado"), ("developer_name", "Asignado a"), ("module", "Módulo"), ("environment", "Ambiente"),
    ("url", "URL"), ("source", "Origen"), ("source_ref", "Referencia"), ("occurrences", "Veces detectada"),
    ("due_date", "Fecha límite"), ("reported_by", "Reportado por"), ("created_at", "Creada"),
    ("updated_at", "Actualizada"), ("description", "Descripción"), ("steps", "Pasos"),
    ("expected", "Resultado esperado"), ("actual", "Resultado actual"),
]


@api.get("/incidents/export")
@token_required
def export_incidents():
    fmt = request.args.get("format", "excel")
    filters = {k: v for k, v in request.args.items() if k != "format"}
    incidents = store.list_incidents(filters)
    rows = [[(i.get(k) if i.get(k) is not None else "") for k, _ in INCIDENT_EXPORT_COLUMNS] for i in incidents]
    headers = [label for _, label in INCIDENT_EXPORT_COLUMNS]
    if fmt == "csv":
        text = io.StringIO()
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow(headers)
        writer.writerows(rows)
        buffer = io.BytesIO(text.getvalue().encode("utf-8-sig"))
        return send_file(buffer, mimetype="text/csv; charset=utf-8", as_attachment=True,
                         download_name="novedades.csv", max_age=0)
    if fmt != "excel":
        raise ValueError("Formato no soportado. Use excel o csv.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Novedades"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for row in rows:
        ws.append(row)
    widths = [11, 50, 11, 11, 13, 22, 18, 11, 40, 18, 14, 10, 12, 18, 19, 19, 50, 40, 30, 50]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="novedades.xlsx", max_age=0)


# ----------------------------------------------------------------------------- configuración
@api.get("/settings")
@token_required
def get_settings():
    return jsonify(store.get_settings())


@api.put("/settings")
@token_required
def update_settings():
    return jsonify(store.update_settings(body()))


@api.post("/settings/test-webhook")
@token_required
def test_webhook():
    try:
        send_webhook(store.get_settings(include_secrets=True), "Prueba de QA Center",
                     "Si ves este mensaje, las novedades llegarán a este canal.")
    except NotificationError as exc:
        return jsonify({"error": str(exc)}), 502
    return jsonify({"message": "Mensaje de prueba enviado"})


@api.post("/settings/test-email")
@token_required
def test_email():
    to = str(body().get("to") or "").strip()
    if not to:
        raise ValueError("Indica el correo de destino")
    try:
        send_email(store.get_settings(include_secrets=True), to, "Prueba de QA Center",
                   "Si recibes este correo, las novedades se enviarán correctamente a los programadores.")
    except NotificationError as exc:
        return jsonify({"error": str(exc)}), 502
    return jsonify({"message": f"Correo de prueba enviado a {to}"})
