"""Registro persistente de ejecuciones (SQLite), cálculo de métricas y generación de reportes (PDF, Excel, CSV)."""
import csv
import datetime as dt
import io
import os
import sqlite3
import threading
from contextlib import contextmanager

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ENVIRONMENTS = ["Local", "QA", "Staging", "Producción"]
TRIGGER_SCHEDULED = "Programada"
TRIGGER_MANUAL = "Manual"
TRIGGER_REGISTERED = "Registro manual"
TRIGGER_IMPORTED = "Importación"
TOOLS = ["Postman", "Selenium", "Cypress", "Playwright", "JMeter", "Pytest", "Manual", "Otro"]
SYSTEM_TOOL = "Sistema"

DEFAULT_DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "qa_reports.db")

REPORT_COLUMNS = [
    ("run_id", "ID Ejecución"),
    ("suite", "Suite"),
    ("tool", "Herramienta"),
    ("environment", "Ambiente"),
    ("status", "Estado"),
    ("total", "Total"),
    ("passed", "Exitosas"),
    ("failed", "Fallidas"),
    ("skipped", "Omitidas"),
    ("pass_rate", "Pass Rate (%)"),
    ("avg_response_ms", "Resp. Prom. (ms)"),
    ("duration_s", "Duración (s)"),
    ("executed_by", "Ejecutado por"),
    ("trigger", "Origen"),
    ("executed_at", "Fecha de ejecución"),
]

EXPORT_FORMATS = {
    "csv": ("text/csv; charset=utf-8", "reporte_pruebas.csv"),
    "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "reporte_pruebas.xlsx"),
    "pdf": ("application/pdf", "reporte_pruebas.pdf"),
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS test_runs (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id          TEXT UNIQUE,
    batch_id        TEXT,
    suite           TEXT NOT NULL,
    environment     TEXT NOT NULL,
    status          TEXT NOT NULL,
    total           INTEGER NOT NULL,
    passed          INTEGER NOT NULL,
    failed          INTEGER NOT NULL,
    skipped         INTEGER NOT NULL,
    pass_rate       REAL NOT NULL,
    avg_response_ms INTEGER NOT NULL,
    duration_s      REAL NOT NULL,
    executed_by     TEXT NOT NULL,
    "trigger"       TEXT NOT NULL,
    executed_at     TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS test_cases (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id      TEXT NOT NULL,
    batch_id    TEXT NOT NULL,
    suite       TEXT NOT NULL,
    name        TEXT NOT NULL,
    status      TEXT NOT NULL,
    response_ms REAL NOT NULL,
    detail      TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS idx_test_cases_batch ON test_cases (batch_id);
CREATE INDEX IF NOT EXISTS idx_test_cases_run ON test_cases (run_id);
"""

# Columnas agregadas después de la primera versión (se crean al abrir una base antigua).
EXTRA_COLUMNS = {
    "tool": "TEXT NOT NULL DEFAULT ''",
    "module": "TEXT NOT NULL DEFAULT ''",
    "url": "TEXT NOT NULL DEFAULT ''",
    "notes": "TEXT NOT NULL DEFAULT ''",
    "evidence_url": "TEXT NOT NULL DEFAULT ''",
    "updated_at": "TEXT",
}

RUN_FIELDS = ([key for key, _ in REPORT_COLUMNS if key != "run_id"] + ["batch_id"]
              + [c for c in EXTRA_COLUMNS if c not in ("tool", "updated_at")])
EDITABLE_FIELDS = [f for f in RUN_FIELDS if f not in ("batch_id", "trigger", "executed_by")] + ["updated_at"]


def now_str() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def build_run_record(suite, environment, total, failed, skipped, avg_response_ms, duration_s,
                     executed_by, trigger, batch_id=None, tool=SYSTEM_TOOL, module="", url="", notes="",
                     evidence_url="", executed_at=None) -> dict:
    """Arma un registro de ejecución calculando exitosas, estado y pass rate."""
    passed = total - failed - skipped
    return {
        "batch_id": batch_id,
        "suite": suite,
        "tool": tool,
        "environment": environment,
        "status": "Passed" if failed == 0 else "Failed",
        "total": total,
        "passed": passed,
        "failed": failed,
        "skipped": skipped,
        "pass_rate": round(passed / total * 100, 2) if total else 0.0,
        "avg_response_ms": int(round(avg_response_ms)),
        "duration_s": round(float(duration_s), 2),
        "executed_by": executed_by,
        "trigger": trigger,
        "executed_at": executed_at or now_str(),
        "module": module,
        "url": url,
        "notes": notes,
        "evidence_url": evidence_url,
    }


class TestRunRepository:
    """Almacén persistente (SQLite) de ejecuciones y casos de prueba."""

    __test__ = False  # evita que pytest intente recolectarla como clase de pruebas

    def __init__(self, db_path: str):
        self.db_path = db_path
        self._lock = threading.Lock()
        os.makedirs(os.path.dirname(os.path.abspath(db_path)), exist_ok=True)
        with self._connection() as conn:
            conn.executescript(SCHEMA)
            existing = {row["name"] for row in conn.execute("PRAGMA table_info(test_runs)")}
            for column, ddl in EXTRA_COLUMNS.items():
                if column not in existing:
                    conn.execute(f'ALTER TABLE test_runs ADD COLUMN "{column}" {ddl}')

    @contextmanager
    def _connection(self):
        conn = sqlite3.connect(self.db_path, timeout=10)
        conn.row_factory = sqlite3.Row
        try:
            yield conn
            conn.commit()
        finally:
            conn.close()

    def add_run(self, run: dict, cases=None) -> dict:
        """Inserta una ejecución y sus casos. Los casos de ejecuciones sin lote (importadas)
        se guardan con ``batch_id = run_id`` para no mezclarlos con los lotes automáticos."""
        columns = ", ".join(f'"{f}"' for f in RUN_FIELDS)
        placeholders = ", ".join("?" for _ in RUN_FIELDS)
        with self._lock, self._connection() as conn:
            cursor = conn.execute(
                f"INSERT INTO test_runs ({columns}) VALUES ({placeholders})",
                [run.get(f, "") if f in EXTRA_COLUMNS else run[f] for f in RUN_FIELDS],
            )
            run_id = f"RUN-{cursor.lastrowid + 1000}"
            conn.execute("UPDATE test_runs SET run_id = ? WHERE id = ?", (run_id, cursor.lastrowid))
            conn.executemany(
                "INSERT INTO test_cases (run_id, batch_id, suite, name, status, response_ms, detail) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                [(run_id, run["batch_id"] or run_id, run["suite"], c["name"], c["status"], c["response_ms"],
                  c["detail"]) for c in cases or []],
            )
        return {**run, "run_id": run_id}

    def _select_runs(self, where="", params=()):
        fields = ", ".join(f'"{f}"' for f in RUN_FIELDS + ["updated_at"])
        with self._connection() as conn:
            rows = conn.execute(f"SELECT run_id, {fields} FROM test_runs {where} ORDER BY id DESC", params).fetchall()
        return [dict(row) for row in rows]

    def get_runs(self) -> list:
        """Ejecuciones ordenadas de la más reciente a la más antigua."""
        return self._select_runs()

    def get_run(self, run_id: str):
        rows = self._select_runs("WHERE run_id = ?", (run_id,))
        return rows[0] if rows else None

    def get_run_cases(self, run_id: str) -> list:
        with self._connection() as conn:
            rows = conn.execute("SELECT id, run_id, suite, name, status, response_ms, detail FROM test_cases "
                                "WHERE run_id = ? ORDER BY id", (run_id,)).fetchall()
        return [dict(r) for r in rows]

    def update_run(self, run_id: str, record: dict) -> dict:
        fields = [f for f in EDITABLE_FIELDS if f in record]
        assignments = ", ".join(f'"{f}" = ?' for f in fields)
        with self._lock, self._connection() as conn:
            conn.execute(f"UPDATE test_runs SET {assignments} WHERE run_id = ?",
                         [record[f] for f in fields] + [run_id])
        return self.get_run(run_id)

    def delete_run(self, run_id: str) -> bool:
        with self._lock, self._connection() as conn:
            deleted = conn.execute("DELETE FROM test_runs WHERE run_id = ?", (run_id,)).rowcount
            conn.execute("DELETE FROM test_cases WHERE run_id = ?", (run_id,))
        return bool(deleted)

    def get_latest_batch(self):
        """Último lote de pruebas automáticas con el detalle de cada caso, o ``None``."""
        with self._connection() as conn:
            head = conn.execute(
                'SELECT batch_id, executed_at, executed_by, "trigger", environment FROM test_runs '
                "WHERE batch_id IS NOT NULL ORDER BY id DESC LIMIT 1"
            ).fetchone()
            if head is None:
                return None
            cases = conn.execute(
                "SELECT run_id, suite, name, status, response_ms, detail FROM test_cases "
                "WHERE batch_id = ? ORDER BY id",
                (head["batch_id"],),
            ).fetchall()
        cases = [dict(c) for c in cases]
        failed = sum(1 for c in cases if c["status"] == "Failed")
        return {
            **dict(head),
            "total": len(cases),
            "passed": len(cases) - failed,
            "failed": failed,
            "cases": cases,
        }


repository = TestRunRepository(os.environ.get("QA_DB_PATH", DEFAULT_DB_PATH))


# ----------------------------------------------------------------------------- Registro manual
def _int_field(data: dict, key: str, label: str, default=None) -> int:
    value = data.get(key, default)
    if value is None or value == "":
        raise ValueError(f"El campo '{label}' es obligatorio")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"El campo '{label}' debe ser numérico") from None
    if number < 0 or number != int(number):
        raise ValueError(f"El campo '{label}' debe ser un entero mayor o igual a 0")
    return int(number)


def _text_field(data: dict, key: str, max_len: int) -> str:
    value = str(data.get(key) or "").strip()
    if len(value) > max_len:
        raise ValueError(f"El campo '{key}' admite máximo {max_len} caracteres")
    return value


def parse_datetime(value) -> str:
    """Acepta 'YYYY-MM-DD HH:MM[:SS]' o el formato de <input type=datetime-local>. Vacío = ahora."""
    text = str(value or "").strip().replace("T", " ")
    if not text:
        return now_str()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    raise ValueError("Fecha de ejecución inválida (use AAAA-MM-DD HH:MM)")


def validate_run_payload(data: dict, executed_by: str, trigger=TRIGGER_REGISTERED) -> dict:
    """Valida los datos de una ejecución registrada a mano y devuelve el registro calculado."""
    if not isinstance(data, dict):
        raise ValueError("El cuerpo de la petición debe ser JSON")
    suite = str(data.get("suite") or "").strip()
    if not 1 <= len(suite) <= 40:
        raise ValueError("La suite es obligatoria (máximo 40 caracteres)")
    environment = str(data.get("environment") or "").strip()
    if environment not in ENVIRONMENTS:
        raise ValueError(f"Ambiente inválido. Use: {', '.join(ENVIRONMENTS)}")
    total = _int_field(data, "total", "Total")
    failed = _int_field(data, "failed", "Fallidas", default=0)
    skipped = _int_field(data, "skipped", "Omitidas", default=0)
    avg_response_ms = _int_field(data, "avg_response_ms", "Resp. promedio (ms)", default=0)
    try:
        duration_s = float(data.get("duration_s") or 0)
    except (TypeError, ValueError):
        raise ValueError("El campo 'Duración (s)' debe ser numérico") from None
    if total < 1:
        raise ValueError("El total de pruebas debe ser al menos 1")
    if failed + skipped > total:
        raise ValueError("Fallidas + Omitidas no puede superar el total")
    if duration_s < 0:
        raise ValueError("La duración no puede ser negativa")
    tool = str(data.get("tool") or "Manual").strip()
    if tool not in TOOLS and not (tool == SYSTEM_TOOL and trigger != TRIGGER_REGISTERED):
        raise ValueError(f"Herramienta inválida. Use: {', '.join(TOOLS)}")

    return build_run_record(
        suite, environment, total, failed, skipped, avg_response_ms, duration_s, executed_by, trigger,
        tool=tool, module=_text_field(data, "module", 80), url=_text_field(data, "url", 500),
        notes=_text_field(data, "notes", 4000), evidence_url=_text_field(data, "evidence_url", 500),
        executed_at=parse_datetime(data.get("executed_at")),
    )


def register_manual_run(data: dict, executed_by: str) -> dict:
    """Valida y registra una ejecución ingresada manualmente. Lanza ``ValueError`` si es inválida."""
    return repository.add_run(validate_run_payload(data, executed_by))


def update_manual_run(run_id: str, data: dict):
    """Edita una ejecución. Devuelve ``None`` si no existe; ``ValueError`` si los datos son inválidos."""
    current = repository.get_run(run_id)
    if current is None:
        return None
    merged = {**current, **(data if isinstance(data, dict) else {})}
    record = validate_run_payload(merged, current["executed_by"], current["trigger"])
    record["updated_at"] = now_str()
    return repository.update_run(run_id, record)


# ----------------------------------------------------------------------------- Métricas
def compute_metrics(runs: list) -> dict:
    total_tests = sum(r["total"] for r in runs)
    passed = sum(r["passed"] for r in runs)
    failed = sum(r["failed"] for r in runs)
    skipped = sum(r["skipped"] for r in runs)
    avg_response = round(sum(r["avg_response_ms"] for r in runs) / len(runs), 2) if runs else 0.0

    by_suite = []
    for suite in sorted({r["suite"] for r in runs}):
        suite_runs = [r for r in runs if r["suite"] == suite]
        s_total = sum(r["total"] for r in suite_runs)
        s_passed = sum(r["passed"] for r in suite_runs)
        by_suite.append({
            "suite": suite,
            "executions": len(suite_runs),
            "total": s_total,
            "passed": s_passed,
            "failed": sum(r["failed"] for r in suite_runs),
            "skipped": sum(r["skipped"] for r in suite_runs),
            "pass_rate": round(s_passed / s_total * 100, 2) if s_total else 0.0,
        })

    return {
        "generated_at": now_str(),
        "total_executions": len(runs),
        "total_tests": total_tests,
        "passed": passed,
        "failed": failed,
        "skipped": skipped,
        "pass_rate": round(passed / total_tests * 100, 2) if total_tests else 0.0,
        "bugs_found": failed,
        "avg_response_time_ms": avg_response,
        "by_suite": by_suite,
        "recent_executions": runs,
    }


def get_dashboard_metrics() -> dict:
    metrics = compute_metrics(repository.get_runs())
    metrics["last_batch"] = repository.get_latest_batch()
    return metrics


# ----------------------------------------------------------------------------- CSV
def generate_csv(runs: list) -> io.BytesIO:
    text = io.StringIO()
    writer = csv.writer(text, lineterminator="\r\n")
    writer.writerow([label for _, label in REPORT_COLUMNS])
    for run in runs:
        writer.writerow([run[key] for key, _ in REPORT_COLUMNS])
    # utf-8-sig agrega BOM para que Excel interprete correctamente los acentos.
    return io.BytesIO(text.getvalue().encode("utf-8-sig"))


# ----------------------------------------------------------------------------- Excel
def generate_excel(metrics: dict) -> io.BytesIO:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fail_fill = PatternFill("solid", fgColor="FEE2E2")
    pass_fill = PatternFill("solid", fgColor="DCFCE7")

    # Hoja 1: resumen ejecutivo
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Reporte Ejecutivo de QA"
    ws["A1"].font = Font(bold=True, size=16, color="1E3A8A")
    ws["A2"] = f"Generado: {metrics['generated_at']}"
    ws["A2"].font = Font(italic=True, color="64748B")
    summary = [
        ("Ejecuciones", metrics["total_executions"]),
        ("Pruebas ejecutadas", metrics["total_tests"]),
        ("Exitosas", metrics["passed"]),
        ("Fallidas (bugs detectados)", metrics["failed"]),
        ("Omitidas", metrics["skipped"]),
        ("Pass Rate (%)", metrics["pass_rate"]),
        ("Tiempo de respuesta promedio (ms)", metrics["avg_response_time_ms"]),
    ]
    for col, title in enumerate(("Métrica", "Valor"), start=1):
        cell = ws.cell(row=4, column=col, value=title)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
    for i, (label, value) in enumerate(summary, start=5):
        ws.cell(row=i, column=1, value=label).border = border
        ws.cell(row=i, column=2, value=value).border = border

    start = 5 + len(summary) + 2
    ws.cell(row=start - 1, column=1, value="Resultados por suite").font = Font(bold=True, size=12)
    suite_headers = ["Suite", "Ejecuciones", "Total", "Exitosas", "Fallidas", "Omitidas", "Pass Rate (%)"]
    for col, title in enumerate(suite_headers, start=1):
        cell = ws.cell(row=start, column=col, value=title)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
    for i, s in enumerate(metrics["by_suite"], start=start + 1):
        values = [s["suite"], s["executions"], s["total"], s["passed"], s["failed"], s["skipped"], s["pass_rate"]]
        for col, value in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=value).border = border
    ws.column_dimensions["A"].width = 36
    for col in range(2, len(suite_headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # Hoja 2: detalle de ejecuciones
    ws2 = wb.create_sheet("Ejecuciones")
    for col, (_, label) in enumerate(REPORT_COLUMNS, start=1):
        cell = ws2.cell(row=1, column=col, value=label)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    status_col = [key for key, _ in REPORT_COLUMNS].index("status") + 1
    for row, run in enumerate(metrics["recent_executions"], start=2):
        for col, (key, _) in enumerate(REPORT_COLUMNS, start=1):
            cell = ws2.cell(row=row, column=col, value=run[key])
            cell.border = border
            if col == status_col:
                cell.fill = pass_fill if run["status"] == "Passed" else fail_fill
    widths = [14, 16, 13, 13, 10, 8, 10, 10, 10, 13, 16, 12, 20, 16, 20]
    for col, width in enumerate(widths, start=1):
        ws2.column_dimensions[get_column_letter(col)].width = width
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_COLUMNS))}{len(metrics['recent_executions']) + 1}"

    # Hoja 3: detalle de casos del último lote automático
    batch = metrics.get("last_batch")
    if batch:
        ws3 = wb.create_sheet("Último lote")
        ws3["A1"] = f"Lote {batch['batch_id']} - {batch['executed_at']} ({batch['trigger']})"
        ws3["A1"].font = Font(bold=True, size=12, color="1E3A8A")
        case_headers = ["Ejecución", "Suite", "Caso de prueba", "Estado", "Tiempo (ms)", "Detalle"]
        for col, title in enumerate(case_headers, start=1):
            cell = ws3.cell(row=3, column=col, value=title)
            cell.font, cell.fill, cell.border = header_font, header_fill, border
        for row, case in enumerate(batch["cases"], start=4):
            values = [case["run_id"], case["suite"], case["name"], case["status"], case["response_ms"], case["detail"]]
            for col, value in enumerate(values, start=1):
                cell = ws3.cell(row=row, column=col, value=value)
                cell.border = border
                if col == 4:
                    cell.fill = pass_fill if case["status"] == "Passed" else fail_fill
        for col, width in enumerate([14, 16, 55, 10, 12, 60], start=1):
            ws3.column_dimensions[get_column_letter(col)].width = width

    def table_sheet(title, headers, rows, widths, status_col=None, bad_values=()):
        sheet = wb.create_sheet(title)
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font, cell.fill, cell.border = header_font, header_fill, border
        for r, values in enumerate(rows, start=2):
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 3) and title == "Novedades")
                if status_col and col == status_col:
                    cell.fill = fail_fill if value in bad_values else pass_fill
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # Hoja 4: novedades para programadores
    if metrics.get("incidents_list") is not None:
        table_sheet("Novedades", INCIDENT_REPORT_HEADERS, [incident_row(i) for i in metrics["incidents_list"]],
                    [11, 48, 50, 11, 13, 22, 16, 11, 18, 17], status_col=5,
                    bad_values=("Nueva", "Asignada", "En progreso", "Reabierta"))
    # Hoja 5: análisis de URL (errores de consola y red)
    if metrics.get("scans_list") is not None:
        table_sheet("Análisis URL", SCAN_REPORT_HEADERS, [scan_row(sc) for sc in metrics["scans_list"]],
                    [12, 45, 18, 13, 10, 10, 12, 8, 13, 19], status_col=4, bad_values=("Con errores", "Error"))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ----------------------------------------------------------------------------- PDF
def generate_pdf(metrics: dict) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title="Reporte Ejecutivo de QA", author="QA Center",
    )
    styles = getSampleStyleSheet()
    primary = colors.HexColor("#1E3A8A")
    grid = colors.HexColor("#CBD5E1")
    pass_color = colors.HexColor("#15803D")
    fail_color = colors.HexColor("#B91C1C")

    def styled_table(data, font_size=8, col_widths=None):
        table = Table(data, colWidths=col_widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), primary),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("GRID", (0, 0), (-1, -1), 0.4, grid),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        for i in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F1F5F9")))
        table.setStyle(TableStyle(style))
        return table

    def color_status(table, rows, col):
        for i, status in enumerate(rows, start=1):
            table.setStyle(TableStyle([
                ("TEXTCOLOR", (col, i), (col, i), pass_color if status == "Passed" else fail_color),
                ("FONTNAME", (col, i), (col, i), "Helvetica-Bold"),
            ]))

    story = [
        Paragraph("Reporte Ejecutivo de QA", styles["Title"]),
        Paragraph(f"Generado: {metrics['generated_at']}", styles["Normal"]),
        Spacer(1, 0.5 * cm),
        Paragraph("Resumen", styles["Heading2"]),
        styled_table([
            ["Ejecuciones", "Pruebas", "Exitosas", "Fallidas", "Omitidas", "Pass Rate", "Resp. promedio"],
            [metrics["total_executions"], metrics["total_tests"], metrics["passed"], metrics["failed"],
             metrics["skipped"], f"{metrics['pass_rate']} %", f"{metrics['avg_response_time_ms']} ms"],
        ]),
    ]

    if metrics["by_suite"]:
        story += [
            Spacer(1, 0.5 * cm),
            Paragraph("Resultados por suite", styles["Heading2"]),
            styled_table(
                [["Suite", "Ejecuciones", "Total", "Exitosas", "Fallidas", "Omitidas", "Pass Rate (%)"]]
                + [[s["suite"], s["executions"], s["total"], s["passed"], s["failed"], s["skipped"], s["pass_rate"]]
                   for s in metrics["by_suite"]]
            ),
        ]

    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    batch = metrics.get("last_batch")
    if batch:
        cases_table = styled_table(
            [["Suite", "Caso de prueba", "Estado", "Tiempo (ms)", "Detalle"]]
            + [[c["suite"], Paragraph(c["name"], cell_style), c["status"], c["response_ms"],
                Paragraph(c["detail"] or "-", cell_style)] for c in batch["cases"]],
            font_size=7,
            col_widths=[2.8 * cm, 9 * cm, 1.8 * cm, 2 * cm, 11.5 * cm],
        )
        color_status(cases_table, [c["status"] for c in batch["cases"]], 2)
        story += [
            Spacer(1, 0.5 * cm),
            Paragraph(f"Último lote automático: {batch['batch_id']} ({batch['executed_at']}) — "
                      f"{batch['passed']}/{batch['total']} casos exitosos", styles["Heading2"]),
            cases_table,
        ]

    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    incidents = metrics.get("incidents_list")
    if incidents is not None:
        open_incidents = [i for i in incidents if i["status"] in ("Nueva", "Asignada", "En progreso", "Reabierta")]
        story += [Spacer(1, 0.5 * cm),
                  Paragraph(f"Novedades abiertas: {len(open_incidents)} de {len(incidents)}", styles["Heading2"])]
        if open_incidents:
            story.append(styled_table(
                [["Código", "Título", "Severidad", "Estado", "Asignado a", "Módulo", "Origen", "Creada"]]
                + [[i["code"], Paragraph(i["title"], cell_style), i["severity"], i["status"],
                    i.get("developer_name") or "Sin asignar", i["module"] or "-", i["source"], i["created_at"][:16]]
                   for i in open_incidents],
                font_size=7, col_widths=[1.8 * cm, 9.5 * cm, 1.8 * cm, 2 * cm, 3.4 * cm, 3 * cm, 3 * cm, 2.6 * cm]))
        else:
            story.append(Paragraph("No hay novedades abiertas.", styles["Normal"]))

    scans = metrics.get("scans_list")
    if scans:
        scan_table = styled_table(
            [["ID", "URL", "Aplicación", "Estado", "Errores", "Red", "Advert.", "Fecha"]]
            + [[sc["scan_id"], Paragraph(sc["url"], cell_style), sc.get("target_name") or "-", sc["status"],
                sc["errors"], sc["network_errors"], sc["warnings"], sc["started_at"][:16]] for sc in scans[:30]],
            font_size=7, col_widths=[2 * cm, 9.5 * cm, 3.5 * cm, 2.4 * cm, 1.6 * cm, 1.4 * cm, 1.6 * cm, 2.8 * cm])
        story += [Spacer(1, 0.5 * cm), Paragraph("Análisis de URL recientes (errores de consola y red)",
                                                  styles["Heading2"]), scan_table]

    story += [Spacer(1, 0.5 * cm), Paragraph("Detalle de ejecuciones", styles["Heading2"])]
    if metrics["recent_executions"]:
        detail = styled_table(
            [[label for _, label in REPORT_COLUMNS]]
            + [[run[key] for key, _ in REPORT_COLUMNS] for run in metrics["recent_executions"]],
            font_size=6.5,
        )
        color_status(detail, [r["status"] for r in metrics["recent_executions"]],
                     [key for key, _ in REPORT_COLUMNS].index("status"))
        story.append(detail)
    else:
        story.append(Paragraph("Aún no hay ejecuciones registradas.", styles["Normal"]))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(document.leftMargin, 0.8 * cm, "QA Center - Reporte confidencial")
        canvas.drawRightString(document.pagesize[0] - document.rightMargin, 0.8 * cm, f"Página {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)
    return buffer


INCIDENT_REPORT_HEADERS = ["Código", "Título", "Resultado actual / error", "Severidad", "Estado", "Asignado a",
                           "Módulo", "Ambiente", "Origen", "Creada"]
SCAN_REPORT_HEADERS = ["ID", "URL", "Aplicación", "Estado", "Errores", "Red", "Advertencias", "HTTP", "Carga (ms)",
                       "Fecha"]


def incident_row(i: dict) -> list:
    return [i["code"], i["title"], i["actual"], i["severity"], i["status"], i.get("developer_name") or "Sin asignar",
            i["module"], i["environment"], i["source"], i["created_at"]]


def scan_row(sc: dict) -> list:
    return [sc["scan_id"], sc["url"], sc.get("target_name") or "", sc["status"], sc["errors"], sc["network_errors"],
            sc["warnings"], sc["http_status"], sc["load_ms"], sc["started_at"]]


def export_report(report_format: str, extra=None):
    """Genera el reporte solicitado. Devuelve ``(buffer, mimetype, filename)``.

    ``extra`` agrega secciones opcionales: ``incidents_list`` (novedades) y ``scans_list`` (análisis de URL).
    Lanza ``ValueError`` si el formato no está soportado.
    """
    fmt = (report_format or "").strip().lower()
    if fmt not in EXPORT_FORMATS:
        raise ValueError(f"Formato no soportado: '{report_format}'. Use pdf, excel o csv.")
    metrics = {**get_dashboard_metrics(), **(extra or {})}
    if fmt == "csv":
        buffer = generate_csv(metrics["recent_executions"])
    elif fmt == "excel":
        buffer = generate_excel(metrics)
    else:
        buffer = generate_pdf(metrics)
    mimetype, filename = EXPORT_FORMATS[fmt]
    return buffer, mimetype, filename
