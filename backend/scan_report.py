"""Reporte de un análisis de URL (PDF con captura de pantalla o Excel) para entregar a los programadores."""
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from reports_engine import now_str

TYPE_LABELS = {"console": "Consola", "javascript": "JavaScript", "network": "Red", "navegador": "Navegador"}
LEVEL_LABELS = {"error": "Error", "warning": "Advertencia"}
MIME = {
    "pdf": "application/pdf",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _summary_rows(scan):
    return [
        ("URL", scan["url"]), ("Aplicación", scan.get("target_name") or "-"), ("Estado", scan["status"]),
        ("Motor", scan["engine"] or "-"), ("HTTP", scan["http_status"] or "-"),
        ("Carga (ms)", scan["load_ms"] if scan["load_ms"] is not None else "-"),
        ("Páginas revisadas", scan["pages_scanned"]), ("Errores JS / consola", scan["errors"]),
        ("Errores de red", scan["network_errors"]), ("Advertencias", scan["warnings"]),
        ("Ejecutado por", f"{scan['executed_by']} ({scan['trigger']})"), ("Fecha", scan["started_at"]),
        ("Observación", scan["message"] or "-"),
    ]


def scan_pdf(scan, screenshot_path=None) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm,
                            bottomMargin=1.5 * cm, title=f"Análisis {scan['scan_id']}", author="QA Center")
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    mono = styles["Code"].clone("mono", fontSize=7, leading=8.5)
    primary = colors.HexColor("#1E3A8A")
    grid = colors.HexColor("#CBD5E1")

    summary = Table([[Paragraph(f"<b>{k}</b>", cell), Paragraph(str(v), cell)] for k, v in _summary_rows(scan)],
                    colWidths=[4.2 * cm, 13.8 * cm])
    summary.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, grid), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                 ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9"))]))
    story = [Paragraph(f"Reporte de análisis {scan['scan_id']}", styles["Title"]),
             Paragraph(f"Generado: {now_str()}", styles["Normal"]), Spacer(1, 0.4 * cm), summary]

    issues = scan.get("issues") or []
    story += [Spacer(1, 0.5 * cm), Paragraph(f"Errores detectados ({len(issues)})", styles["Heading2"])]
    if issues:
        rows = [["Nivel", "Tipo", "Mensaje", "Página / recurso", "Novedad"]]
        for i in issues:
            rows.append([LEVEL_LABELS.get(i["level"], i["level"]), TYPE_LABELS.get(i["source"], i["source"]),
                         Paragraph(_escape(i["message"][:1500]).replace("\n", "<br/>"), mono),
                         Paragraph(_escape(i["resource_url"] or i["page_url"]), cell), i.get("incident_code") or "-"])
        table = Table(rows, colWidths=[1.8 * cm, 2 * cm, 8.4 * cm, 4.2 * cm, 1.6 * cm], repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), primary), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                 ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                 ("GRID", (0, 0), (-1, -1), 0.4, grid), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for n, i in enumerate(issues, start=1):
            if i["level"] == "error":
                style.append(("TEXTCOLOR", (0, n), (0, n), colors.HexColor("#B91C1C")))
        table.setStyle(TableStyle(style))
        story.append(table)
    else:
        story.append(Paragraph("No se detectaron errores en la consola ni en la red.", styles["Normal"]))

    if screenshot_path and os.path.isfile(screenshot_path):
        image = Image(screenshot_path)
        ratio = image.imageHeight / float(image.imageWidth)
        image.drawWidth, image.drawHeight = 16 * cm, 16 * cm * ratio
        story += [Spacer(1, 0.5 * cm), Paragraph("Captura de la página", styles["Heading2"]), image]

    doc.build(story)
    buffer.seek(0)
    return buffer


def scan_excel(scan) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Análisis {scan['scan_id']}"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    for row, (key, value) in enumerate(_summary_rows(scan), start=3):
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    ws2 = wb.create_sheet("Errores")
    headers = ["Nivel", "Tipo", "Mensaje", "Página", "Recurso", "Ubicación", "Novedad"]
    ws2.append(headers)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
    for i in scan.get("issues") or []:
        ws2.append([LEVEL_LABELS.get(i["level"], i["level"]), TYPE_LABELS.get(i["source"], i["source"]), i["message"],
                    i["page_url"], i["resource_url"], i["location"], i.get("incident_code") or ""])
    for row in ws2.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in enumerate([12, 12, 80, 40, 40, 30, 11], start=1):
        ws2.column_dimensions[get_column_letter(col)].width = width
    ws2.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _escape(text):
    return str(text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_scan_report(scan, fmt, screenshot_path=None):
    """Devuelve ``(buffer, mimetype, filename)``. Lanza ``ValueError`` con un formato no soportado."""
    if fmt == "pdf":
        buffer = scan_pdf(scan, screenshot_path)
    elif fmt == "excel":
        buffer = scan_excel(scan)
    else:
        raise ValueError("Formato no soportado. Use pdf o excel.")
    extension = "pdf" if fmt == "pdf" else "xlsx"
    return buffer, MIME[fmt], f"analisis_{scan['scan_id']}.{extension}"
