"""Pruebas de los módulos de QA Center: programadores, aplicaciones, resultados externos (CRUD e importación),
análisis de URL, novedades (CRUD, automáticas, sin duplicados, acciones masivas, exportación) y notificaciones.

Ejecutar (desde la raíz del proyecto):
    python -m pytest tests/test_qa_center.py -v
"""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from contextlib import redirect_stdout
from http.server import BaseHTTPRequestHandler, HTTPServer

os.environ.setdefault("QA_DB_PATH", os.path.join(tempfile.mkdtemp(prefix="qa_center_test_"), "qa_reports.db"))
BACKEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

import monitor  # noqa: E402
from app import create_app  # noqa: E402
from importers import parse_report  # noqa: E402

API = "/api/v1"

NEWMAN_REPORT = {
    "collection": {"info": {"name": "API Facturación"}},
    "run": {
        "stats": {"assertions": {"total": 3, "failed": 1}},
        "timings": {"started": 1000, "completed": 3500, "responseAverage": 120},
        "executions": [
            {"item": {"name": "Login"}, "response": {"code": 200, "responseTime": 100},
             "assertions": [{"assertion": "Status 200"}, {"assertion": "Tiene token"}]},
            {"item": {"name": "Crear factura"}, "response": {"code": 500, "responseTime": 140},
             "assertions": [{"assertion": "Status 201", "error": {"message": "expected 500 to equal 201"}}]},
        ],
    },
}

JUNIT_REPORT = b"""<?xml version="1.0" encoding="utf-8"?>
<testsuites name="pytest"><testsuite name="ui" tests="3" time="4.5">
  <testcase classname="tests.test_ui.LoginTest" name="test_login" time="1.2"/>
  <testcase classname="tests.test_ui.LoginTest" name="test_logout" time="0.8">
    <failure message="NoSuchElementException: #btn-logout">traceback</failure>
  </testcase>
  <testcase classname="tests.test_ui.LoginTest" name="test_reset" time="0"><skipped message="pendiente"/></testcase>
</testsuite></testsuites>"""


def fake_scan(issues, status=None):
    def _scan(url, max_pages=1, wait_seconds=2.0, screenshot_path=None):
        result = {"engine": "Falso", "http_status": 200, "load_ms": 12.5, "page_title": "Demo", "pages_scanned": 1,
                  "message": "", "issues": [dict(i, page_url=url) for i in issues]}
        if status:
            result["status"] = status
        return result
    return _scan


class QACenterTestCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        client = cls.app.test_client()
        with redirect_stdout(io.StringIO()):
            token = client.post(f"{API}/auth/login", json={"email": "admin@test.com",
                                                           "password": "SecurePassword123"}).get_json()["token"]
        cls.headers = {"Authorization": f"Bearer {token}"}
        cls._out = redirect_stdout(io.StringIO())
        cls._out.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls._out.__exit__(None, None, None)

    def setUp(self):
        self.client = self.app.test_client()
        self._original_scan = monitor.scan_url

    def tearDown(self):
        monitor.scan_url = self._original_scan

    def call(self, method, path, expected=200, **kwargs):
        response = getattr(self.client, method)(API + path, headers=self.headers, **kwargs)
        if response.status_code != expected:
            self.fail(f"HTTP {response.status_code} != {expected}: {response.get_data()[:500]!r}")
        return response.get_json() if response.is_json else response

    def new_developer(self, email):
        return self.call("post", "/developers", 201, json={"name": "Dev " + email, "email": email, "area": "Backend"})

    # ------------------------------------------------------------------ seguridad
    def test_new_endpoints_require_token(self):
        for path in ("/developers", "/targets", "/scans", "/incidents", "/settings", "/catalogs"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(API + path).status_code, 401)

    # ------------------------------------------------------------------ programadores
    def test_developer_crud(self):
        dev = self.new_developer("ana@empresa.com")
        self.assertEqual(dev["open_incidents"], 0)
        self.call("post", "/developers", 400, json={"name": "Otra", "email": "ana@empresa.com"})
        self.call("post", "/developers", 400, json={"name": "Sin correo", "email": "no-es-correo"})
        updated = self.call("put", f"/developers/{dev['id']}", json={"area": "Frontend", "active": False})
        self.assertEqual(updated["area"], "Frontend")
        self.assertEqual(updated["active"], 0)

        incident = self.call("post", "/incidents", 201, json={"title": "Botón roto", "developer_id": dev["id"]})
        self.assertEqual(incident["status"], "Asignada")
        self.call("delete", f"/developers/{dev['id']}")
        self.call("delete", f"/developers/{dev['id']}", 404)
        self.assertIsNone(self.call("get", f"/incidents/{incident['id']}")["developer_id"])

    # ------------------------------------------------------------------ aplicaciones
    def test_target_crud_and_validation(self):
        target = self.call("post", "/targets", 201, json={"name": "Web", "url": "localhost:3000", "monitor": True,
                                                           "max_pages": 2})
        self.assertEqual(target["url"], "http://localhost:3000")
        self.assertEqual(target["monitor"], 1)
        self.call("post", "/targets", 400, json={"name": "Mala", "url": "ftp://x"})
        self.call("post", "/targets", 400, json={"name": "Muchas", "url": "http://x", "max_pages": 99})
        self.assertEqual(self.call("put", f"/targets/{target['id']}", json={"monitor": False})["monitor"], 0)
        self.call("delete", f"/targets/{target['id']}")
        self.call("get", f"/targets/{target['id']}", 405)  # solo PUT/DELETE por id

    # ------------------------------------------------------------------ resultados externos
    def test_registered_run_edit_and_delete(self):
        run = self.call("post", "/test-runs", 201, json={
            "tool": "Postman", "suite": "API Pagos", "environment": "QA", "total": 20, "failed": 2,
            "notes": "500 al pagar", "url": "http://localhost:8080/api", "executed_at": "2026-10-18T14:30"})
        self.assertEqual(run["tool"], "Postman")
        self.assertEqual(run["executed_at"], "2026-10-18 14:30:00")

        edited = self.call("put", f"/test-runs/{run['run_id']}", json={"failed": 0, "notes": "Corregido"})
        self.assertEqual(edited["status"], "Passed")
        self.assertEqual(edited["passed"], 20)
        self.assertEqual(edited["notes"], "Corregido")
        self.assertIsNotNone(edited["updated_at"])
        self.call("put", f"/test-runs/{run['run_id']}", 400, json={"failed": 50})
        self.call("put", f"/test-runs/{run['run_id']}", 400, json={"tool": "Excel"})

        listed = self.call("get", "/test-runs?origin=external")
        self.assertIn(run["run_id"], [r["run_id"] for r in listed])
        self.call("delete", f"/test-runs/{run['run_id']}")
        self.call("delete", f"/test-runs/{run['run_id']}", 404)

    def test_run_failures_become_incident(self):
        run = self.call("post", "/test-runs", 201, json={"tool": "Selenium", "suite": "Checkout", "environment": "QA",
                                                         "total": 5, "failed": 1})
        created = self.call("post", f"/test-runs/{run['run_id']}/incidents", 201, json={})["created"]
        self.assertEqual(len(created), 1)
        self.assertEqual(created[0]["source_ref"], run["run_id"])

    def test_import_newman_json(self):
        data = {"file": (io.BytesIO(json.dumps(NEWMAN_REPORT).encode()), "resultado.json"),
                "environment": "Staging", "create_incidents": "1"}
        run = self.call("post", "/test-runs/import", 201, data=data, content_type="multipart/form-data")
        self.assertEqual(run["tool"], "Postman")
        self.assertEqual(run["suite"], "API Facturación")
        self.assertEqual((run["total"], run["failed"]), (3, 1))
        self.assertEqual(run["trigger"], "Importación")
        self.assertEqual(len(run["incidents_created"]), 1)
        cases = self.call("get", f"/test-runs/{run['run_id']}")["cases"]
        failed = [c for c in cases if c["status"] == "Failed"]
        self.assertIn("expected 500 to equal 201", failed[0]["detail"])
        # Importar de nuevo el mismo fallo no duplica la novedad: suma una reaparición.
        data["file"] = (io.BytesIO(json.dumps(NEWMAN_REPORT).encode()), "resultado.json")
        again = self.call("post", "/test-runs/import", 201, data=data, content_type="multipart/form-data")
        self.assertEqual(again["incidents_created"], [])

    def test_import_junit_xml(self):
        data = {"file": (io.BytesIO(JUNIT_REPORT), "resultado.xml"), "tool": "Selenium", "environment": "QA"}
        run = self.call("post", "/test-runs/import", 201, data=data, content_type="multipart/form-data")
        self.assertEqual((run["tool"], run["total"], run["failed"], run["skipped"]), ("Selenium", 3, 1, 1))

    def test_import_rejects_postman_collection_and_garbage(self):
        collection = {"info": {"name": "x"}, "item": []}
        with self.assertRaisesRegex(ValueError, "Newman"):
            parse_report("c.json", json.dumps(collection).encode())
        with self.assertRaisesRegex(ValueError, "no reconocido"):
            parse_report("x.txt", b"hola")
        data = {"file": (io.BytesIO(b"<html></html>"), "x.xml")}
        self.call("post", "/test-runs/import", 400, data=data, content_type="multipart/form-data")

    # ------------------------------------------------------------------ análisis de URL
    def test_scan_creates_incidents_without_duplicates(self):
        dev = self.new_developer("front@empresa.com")
        target = self.call("post", "/targets", 201, json={"name": "Portal", "url": "http://localhost:4000",
                                                           "developer_id": dev["id"]})
        issues = [
            {"level": "error", "source": "javascript", "message": "TypeError: x is undefined at app.js:10:5"},
            {"level": "error", "source": "network", "message": "HTTP 404 Not Found en Script",
             "resource_url": "http://localhost:4000/app.js"},
            {"level": "warning", "source": "console", "message": "API obsoleta"},
        ]
        monitor.scan_url = fake_scan(issues)
        scan = monitor.start_scan(target["url"], "Tester", target_id=target["id"], auto_incidents=True, wait=True)
        self.assertEqual((scan["status"], scan["errors"], scan["network_errors"], scan["warnings"]),
                         ("Con errores", 1, 1, 1))

        detail = self.call("get", f"/scans/{scan['scan_id']}")
        linked = [i for i in detail["issues"] if i["incident_id"]]
        self.assertEqual(len(linked), 2, "Solo los errores (no advertencias) generan novedades por defecto")
        incident = self.call("get", f"/incidents/{linked[0]['incident_id']}")
        self.assertEqual(incident["developer_id"], dev["id"], "Se asigna al responsable de la aplicación")
        self.assertEqual(incident["source"], "Análisis de URL")
        self.assertIn("F12", incident["steps"])

        # Mismo error en otro análisis (con distinto número de línea): no se duplica.
        issues[0]["message"] = "TypeError: x is undefined at app.js:11:7"
        monitor.scan_url = fake_scan(issues)
        second = monitor.start_scan(target["url"], "Tester", target_id=target["id"], auto_incidents=True, wait=True)
        incidents = self.call("get", f"/incidents?q={target['name']}")
        self.assertEqual(len(incidents), 2)
        self.assertTrue(all(i["occurrences"] == 2 for i in incidents))

        # La advertencia se puede pasar a mano.
        warning = [i for i in self.call("get", f"/scans/{second['scan_id']}")["issues"] if i["level"] == "warning"][0]
        summary = self.call("post", f"/scans/{second['scan_id']}/incidents", 201, json={"issue_ids": [warning["id"]]})
        self.assertEqual(len(summary["created"]), 1)
        self.assertEqual(summary["created"][0]["severity"], "Baja")

        self.call("delete", f"/scans/{second['scan_id']}")
        self.call("get", f"/scans/{second['scan_id']}", 404)

    def test_scan_unreachable_url_is_reported(self):
        monitor.scan_url = fake_scan([{"level": "error", "source": "network", "message": "No hay conexión"}],
                                     status="Error")
        scan = monitor.start_scan("http://localhost:1", "Tester", auto_incidents=True, wait=True)
        self.assertEqual(scan["status"], "Error")
        self.assertEqual(self.call("get", "/incidents?q=localhost:1"), [], "Un análisis fallido no crea novedades")

    def test_scan_endpoint_validates_input(self):
        self.call("post", "/scans", 400, json={"url": ""})
        self.call("post", "/scans", 400, json={"url": "javascript:alert(1)"})
        self.call("post", "/scans", 400, json={"url": "http://localhost:3000", "max_pages": 50})

    # ------------------------------------------------------------------ novedades
    def test_incident_crud_history_and_comments(self):
        incident = self.call("post", "/incidents", 201, json={
            "title": "Total de factura incorrecto", "severity": "Alta", "priority": "Urgente",
            "steps": "1. Crear factura", "expected": "Total 100", "actual": "Total 90", "url": "localhost:3000/f"})
        self.assertTrue(incident["code"].startswith("NOV-"))
        self.assertEqual(incident["url"], "http://localhost:3000/f")
        self.call("post", "/incidents", 400, json={"title": ""})
        self.call("post", "/incidents", 400, json={"title": "x", "severity": "Gravísima"})

        updated = self.call("put", f"/incidents/{incident['id']}", json={"status": "Resuelta"})
        self.assertEqual(updated["status"], "Resuelta")
        self.call("post", f"/incidents/{incident['id']}/comments", 201, json={"text": "Corregido en v1.2"})
        detail = self.call("get", f"/incidents/{incident['id']}")
        kinds = [h["kind"] for h in detail["history"]]
        self.assertIn("Comentario", kinds)
        self.assertIn("Estado: Nueva → Resuelta", " ".join(h["detail"] for h in detail["history"]))
        self.assertIn("Total de factura incorrecto", detail["share"]["message"])
        self.assertIn("Pasos para reproducir", detail["share"]["message"])

        self.call("delete", f"/incidents/{incident['id']}")
        self.call("get", f"/incidents/{incident['id']}", 404)

    def test_bulk_actions_and_export(self):
        dev = self.new_developer("bulk@empresa.com")
        ids = [self.call("post", "/incidents", 201, json={"title": f"Bulk {n}"})["id"] for n in range(3)]
        self.call("post", "/incidents/bulk", json={"ids": ids, "action": "assign", "developer_id": dev["id"]})
        self.assertTrue(all(self.call("get", f"/incidents/{i}")["developer_id"] == dev["id"] for i in ids))
        self.call("post", "/incidents/bulk", json={"ids": ids[:1], "action": "status", "status": "Cerrada"})
        open_ids = [i["id"] for i in self.call("get", "/incidents?open=1")]
        self.assertNotIn(ids[0], open_ids)
        self.assertIn(ids[1], open_ids)

        csv_body = self.call("get", f"/incidents/export?format=csv&developer_id={dev['id']}").data.decode("utf-8-sig")
        self.assertTrue(csv_body.startswith("Código,Título"))
        self.assertEqual(csv_body.count("Bulk "), 3)
        xlsx = self.call("get", "/incidents/export?format=excel").data
        self.assertEqual(xlsx[:2], b"PK")

        self.call("post", "/incidents/bulk", json={"ids": ids, "action": "delete"})
        self.call("get", f"/incidents/{ids[0]}", 404)

    def test_notify_without_channels_and_manual_share(self):
        incident = self.call("post", "/incidents", 201, json={"title": "Sin canales"})
        self.call("post", f"/incidents/{incident['id']}/notify", 400, json={"channels": []})
        self.call("post", f"/incidents/{incident['id']}/notify", json={"manual_channel": "texto copiado"})
        detail = self.call("get", f"/incidents/{incident['id']}")
        self.assertIsNotNone(detail["notified_at"])
        self.assertEqual(detail["history"][0]["kind"], "Notificación")

    def test_notify_by_webhook(self):
        received = []

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                received.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                self.send_response(200)
                self.end_headers()

            def log_message(self, *args):
                pass

        server = HTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            self.call("put", "/settings", json={"webhook_url": f"http://127.0.0.1:{server.server_port}/hook"})
            self.call("post", "/settings/test-webhook")
            dev = self.new_developer("hook@empresa.com")
            incident = self.call("post", "/incidents", 201, json={"title": "Error 500 en pagos",
                                                                  "developer_id": dev["id"]})
            result = self.call("post", f"/incidents/{incident['id']}/notify",
                               json={"channels": ["webhook"], "note": "Urgente para hoy"})
            self.assertEqual(result["results"], {"webhook": "ok"})
            self.assertIn("Error 500 en pagos", received[-1]["text"])
            self.assertIn("Urgente para hoy", received[-1]["text"])
            self.assertIn("content", received[-1])  # compatible con Discord
            # Sin SMTP configurado el correo falla con un mensaje claro (502) y queda en el historial.
            response = self.client.post(f"{API}/incidents/{incident['id']}/notify", headers=self.headers,
                                        json={"channels": ["email"]})
            self.assertEqual(response.status_code, 502)
            self.assertIn("SMTP", response.get_json()["results"]["email"])
        finally:
            server.shutdown()
            self.call("put", "/settings", json={"webhook_url": ""})

    # ------------------------------------------------------------------ configuración y métricas
    def test_settings_validation_and_secret_masking(self):
        self.call("put", "/settings", 400, json={"monitor_interval_minutes": 0})
        self.call("put", "/settings", 400, json={"webhook_url": "no es url"})
        saved = self.call("put", "/settings", json={"smtp_password": "secreto", "monitor_interval_minutes": 15})
        self.assertEqual(saved["smtp_password"], "")
        self.assertTrue(saved["smtp_password_set"])
        self.assertEqual(saved["monitor_interval_minutes"], "15")
        # Enviar la contraseña vacía conserva la guardada.
        self.assertTrue(self.call("put", "/settings", json={"smtp_password": ""})["smtp_password_set"])

    def test_metrics_include_incidents_and_scans(self):
        metrics = self.call("get", "/dashboard/metrics")
        self.assertIn("open", metrics["incidents"])
        self.assertIn("by_status", metrics["incidents"])
        self.assertIn("total", metrics["scans"])

    # ------------------------------------------------------------------ reportes
    def test_scan_report_pdf_and_excel(self):
        monitor.scan_url = fake_scan([{"level": "error", "source": "console", "message": "Error <b>de</b> prueba"}])
        scan = monitor.start_scan("http://localhost:5555", "Tester", auto_incidents=True, wait=True)
        pdf = self.call("get", f"/scans/{scan['scan_id']}/report?format=pdf")
        self.assertEqual(pdf.data[:4], b"%PDF")
        self.assertIn(f"analisis_{scan['scan_id']}.pdf", pdf.headers["Content-Disposition"])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.call("get", f"/scans/{scan['scan_id']}/report?format=excel").data))
        errors = [row for row in wb["Errores"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(errors[0][2], "Error <b>de</b> prueba")
        self.assertTrue(errors[0][6].startswith("NOV-"))
        self.call("get", f"/scans/{scan['scan_id']}/report?format=docx", 400)

    def test_executive_report_includes_incidents_and_scans(self):
        self.call("post", "/incidents", 201, json={"title": "Para el reporte ejecutivo"})
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.call("get", "/reports/export?format=excel").data))
        self.assertIn("Novedades", wb.sheetnames)
        self.assertIn("Análisis URL", wb.sheetnames)
        titles = [row[1] for row in wb["Novedades"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("Para el reporte ejecutivo", titles)
        self.assertEqual(self.call("get", "/reports/export?format=pdf").data[:4], b"%PDF")

    def test_automatic_report_generation(self):
        self.call("put", "/settings", 400, json={"report_hour": "25"})
        self.call("put", "/settings", 400, json={"report_recipients": "jefe@empresa.com, malo"})
        self.call("put", "/settings", json={"report_hour": "0", "report_last_date": ""})
        self.assertTrue(monitor.daily_report_due())

        result = self.call("post", "/reports/auto", 201)
        self.assertEqual(len(result["files"]), 2)
        self.assertFalse(monitor.daily_report_due(), "No se repite el mismo día")
        listing = self.call("get", "/reports/auto")
        names = [r["name"] for r in listing["reports"]]
        self.assertIn(result["files"][0], names)
        pdf_name = next(n for n in result["files"] if n.endswith(".pdf"))
        self.assertEqual(self.call("get", f"/reports/auto/{pdf_name}").data[:4], b"%PDF")
        self.call("get", "/reports/auto/..%2Fqa_reports.db", 404)
        self.call("delete", f"/reports/auto/{pdf_name}")
        self.call("get", f"/reports/auto/{pdf_name}", 404)

        # Con destinatarios pero sin SMTP, el reporte igual se guarda y se informa el error de envío.
        self.call("put", "/settings", json={"report_recipients": "jefe@empresa.com"})
        result = self.call("post", "/reports/auto", 201)
        self.assertIn("SMTP", result["email_error"])
        self.call("put", "/settings", json={"report_recipients": "", "report_hour": "18"})

    def test_dashboard_has_all_menus(self):
        html = self.client.get("/dashboard").get_data(as_text=True)
        for view in ("resumen", "analizador", "resultados", "automaticas", "novedades", "programadores",
                     "aplicaciones", "reportes", "configuracion"):
            self.assertIn(f'id="view-{view}"', html)
            self.assertIn(f'href="#{view}"', html)


if __name__ == "__main__":
    unittest.main(verbosity=2)
