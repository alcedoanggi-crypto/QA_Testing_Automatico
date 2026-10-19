"""Pruebas de Documentación QA: plan de pruebas, requisitos, casos de prueba (ejecución, defectos y evidencias),
matriz de trazabilidad y reporte de defectos.

Ejecutar (desde la raíz del proyecto):
    python -m pytest tests/test_testdocs.py -v
"""
import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout

os.environ.setdefault("QA_DB_PATH", os.path.join(tempfile.mkdtemp(prefix="qa_docs_test_"), "qa_reports.db"))
BACKEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from openpyxl import load_workbook  # noqa: E402

from app import create_app  # noqa: E402

API = "/api/v1"
def _png():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (40, 20), (220, 38, 38)).save(buffer, format="PNG")
    return buffer.getvalue()


PNG = _png()


class TestDocsCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls._out = redirect_stdout(io.StringIO())
        cls._out.__enter__()
        token = cls.app.test_client().post(f"{API}/auth/login", json={
            "email": "admin@test.com", "password": "SecurePassword123"}).get_json()["token"]
        cls.headers = {"Authorization": f"Bearer {token}"}

    @classmethod
    def tearDownClass(cls):
        cls._out.__exit__(None, None, None)

    def setUp(self):
        self.client = self.app.test_client()

    def call(self, method, path, expected=200, **kwargs):
        response = getattr(self.client, method)(API + path, headers=self.headers, **kwargs)
        if response.status_code != expected:
            self.fail(f"{method.upper()} {path}: HTTP {response.status_code} != {expected}: "
                      f"{response.get_data()[:400]!r}")
        return response.get_json() if response.is_json else response

    def new_plan(self, name="Plan Facturación"):
        return self.call("post", "/plans", 201, json={
            "name": name, "version": "1.0", "responsible": "Anggie", "start_date": "2026-10-01",
            "end_date": "2026-10-31", "objective": "Validar facturación", "scope_in": "Emisión de facturas",
            "features": "Crear, anular y consultar facturas", "acceptance_criteria": "95 % de casos aprobados",
            "resources": "2 testers, ambiente QA", "schedule": "Semana 1: diseño\nSemana 2: ejecución"})

    def test_endpoints_require_token(self):
        for path in ("/plans", "/requirements", "/test-cases", "/rtm", "/docs/summary"):
            self.assertEqual(self.client.get(API + path).status_code, 401, path)

    def test_plan_crud_and_validation(self):
        plan = self.new_plan()
        self.assertTrue(plan["code"].startswith("PLN-"))
        self.assertEqual(plan["summary"]["cases"], 0)
        self.call("post", "/plans", 400, json={"name": ""})
        self.call("post", "/plans", 400, json={"name": "X", "start_date": "2026-10-10", "end_date": "2026-10-01"})
        self.call("post", "/plans", 400, json={"name": "X", "status": "Inventado"})
        updated = self.call("put", f"/plans/{plan['id']}", json={"status": "Aprobado", "version": "1.1"})
        self.assertEqual((updated["status"], updated["version"]), ("Aprobado", "1.1"))
        self.call("delete", f"/plans/{plan['id']}")
        self.call("get", f"/plans/{plan['id']}", 404)

    def test_requirements_crud(self):
        req = self.call("post", "/requirements", 201, json={"code": "hu-01", "title": "Crear factura",
                                                             "kind": "Historia de usuario"})
        self.assertEqual(req["code"], "HU-01")
        auto = self.call("post", "/requirements", 201, json={"title": "Sin código"})
        self.assertTrue(auto["code"].startswith("REQ-"))
        self.call("post", "/requirements", 400, json={"code": "HU-01", "title": "Duplicado"})
        self.call("post", "/requirements", 400, json={"code": "con espacio", "title": "x"})
        self.assertEqual(self.call("put", f"/requirements/{req['id']}", json={"priority": "Alta"})["priority"], "Alta")
        self.call("delete", f"/requirements/{auto['id']}")
        self.call("delete", f"/requirements/{auto['id']}", 404)

    def test_case_lifecycle_execution_defect_and_evidence(self):
        plan = self.new_plan("Plan ciclo")
        req = self.call("post", "/requirements", 201, json={"code": "HU-CICLO", "title": "Login",
                                                             "plan_id": plan["id"]})
        spec = self.call("post", "/test-cases", 201, json={
            "plan_id": plan["id"], "title": "Login con clave incorrecta", "description": "Validar mensaje de error",
            "preconditions": "Usuario registrado", "input_data": "admin@test.com / clave123",
            "steps": "1. Abrir /login\n2. Ingresar datos\n3. Pulsar Ingresar",
            "expected": "Mensaje «Credenciales inválidas»", "requirement_ids": [req["id"]], "type": "Funcional"})
        self.assertTrue(spec["code"].startswith("CP-"))
        self.assertEqual(spec["status"], "No ejecutado")
        self.assertEqual([r["code"] for r in spec["requirements"]], ["HU-CICLO"])
        self.call("post", "/test-cases", 404, json={"title": "x", "requirement_ids": [99999]})
        self.call("post", "/test-cases", 400, json={"title": "x", "type": "Mágico"})

        rtm = self.call("get", f"/rtm?plan_id={plan['id']}")
        self.assertEqual(rtm["rows"][0]["coverage"], "Pendiente")

        dev = self.call("post", "/developers", 201, json={"name": "Dev Docs", "email": "docs@empresa.com"})
        result = self.call("post", f"/test-cases/{spec['id']}/execute", json={
            "status": "Fallido", "actual": "Pantalla en blanco", "create_defect": True, "severity": "Alta",
            "developer_id": dev["id"]})
        self.assertEqual(result["status"], "Fallido")
        defect = result["defect"]
        self.assertEqual(defect["source"], "Caso de prueba")
        self.assertEqual(defect["developer_id"], dev["id"])
        self.assertIn("Precondiciones: Usuario registrado", defect["steps"])
        self.assertEqual(defect["actual"], "Pantalla en blanco")
        self.assertEqual(self.call("get", f"/rtm?plan_id={plan['id']}")["rows"][0]["coverage"], "Con fallas")

        # Volver a fallar no duplica el defecto abierto.
        again = self.call("post", f"/test-cases/{spec['id']}/execute", json={
            "status": "Fallido", "actual": "Sigue en blanco", "create_defect": True})
        self.assertEqual(again["defect"]["id"], defect["id"])
        self.assertEqual(again["defect"]["occurrences"], 2)

        ok = self.call("post", f"/test-cases/{spec['id']}/execute", json={"status": "Aprobado", "actual": "OK"})
        self.assertEqual(len(ok["executions"]), 3)
        self.assertEqual(self.call("get", f"/rtm?plan_id={plan['id']}")["rows"][0]["coverage"], "Validado")
        summary = self.call("get", f"/plans/{plan['id']}")["summary"]
        self.assertEqual((summary["cases"], summary["executed"], summary["pass_rate"], summary["coverage"]),
                         (1, 1, 100.0, 100.0))

        # Evidencias
        uploaded = self.call("post", f"/test-cases/{spec['id']}/evidence", 201, content_type="multipart/form-data",
                             data={"file": [(io.BytesIO(PNG), "captura.png"), (io.BytesIO(b"log"), "consola.log")]})
        self.assertEqual(len(uploaded), 2)
        self.call("post", f"/test-cases/{spec['id']}/evidence", 400, content_type="multipart/form-data",
                  data={"file": (io.BytesIO(b"MZ"), "virus.exe")})
        image = self.call("get", f"/evidence/{uploaded[0]['id']}?inline=1")
        self.assertEqual(image.data, PNG)
        self.assertEqual(self.call("get", f"/test-cases/{spec['id']}")["evidence_count"], 2)

        # Documentos
        pdf = self.call("get", f"/test-cases/{spec['id']}/report?format=pdf")
        self.assertEqual(pdf.data[:4], b"%PDF")
        self.assertEqual(self.call("get", f"/plans/{plan['id']}/report?format=pdf").data[:4], b"%PDF")
        wb = load_workbook(io.BytesIO(self.call("get", f"/plans/{plan['id']}/report?format=excel").data))
        self.assertEqual(wb.sheetnames, ["Plan", "Casos de prueba", "Trazabilidad"])
        self.assertEqual(self.call("get", f"/test-cases/report?format=pdf&plan_id={plan['id']}").data[:4], b"%PDF")
        cases_xlsx = load_workbook(io.BytesIO(self.call("get", "/test-cases/report?format=excel").data))
        headers = [c.value for c in cases_xlsx.active[3]]
        for column in ("Precondiciones", "Datos de entrada", "Pasos", "Resultado esperado", "Resultado obtenido",
                       "Evidencias"):
            self.assertIn(column, headers)

        copy = self.call("post", f"/test-cases/{spec['id']}/duplicate", 201)
        self.assertEqual(copy["status"], "No ejecutado")
        self.assertEqual([r["code"] for r in copy["requirements"]], ["HU-CICLO"])

        self.call("delete", f"/evidence/{uploaded[1]['id']}")
        self.call("delete", f"/test-cases/{spec['id']}")
        self.call("get", f"/test-cases/{spec['id']}", 404)
        self.call("get", f"/evidence/{uploaded[0]['id']}", 404)

    def test_rtm_report_and_uncovered_requirements(self):
        plan = self.new_plan("Plan RTM")
        covered = self.call("post", "/requirements", 201, json={"code": "RTM-1", "title": "A", "plan_id": plan["id"]})
        self.call("post", "/requirements", 201, json={"code": "RTM-2", "title": "B", "plan_id": plan["id"]})
        self.call("post", "/test-cases", 201, json={"title": "Caso A", "plan_id": plan["id"],
                                                    "requirement_ids": [covered["id"]]})
        self.call("post", "/test-cases", 201, json={"title": "Caso huérfano", "plan_id": plan["id"]})
        rtm = self.call("get", f"/rtm?plan_id={plan['id']}")
        self.assertEqual(rtm["totals"]["requirements"], 2)
        self.assertEqual(rtm["totals"]["coverage"], 50.0)
        self.assertEqual({r["code"]: r["coverage"] for r in rtm["rows"]}, {"RTM-1": "Pendiente", "RTM-2": "Sin cobertura"})
        self.assertEqual([c["title"] for c in rtm["orphan_cases"]], ["Caso huérfano"])
        self.assertEqual(self.call("get", f"/rtm/report?format=pdf&plan_id={plan['id']}").data[:4], b"%PDF")
        wb = load_workbook(io.BytesIO(self.call("get", f"/rtm/report?format=excel&plan_id={plan['id']}").data))
        self.assertIn("Matriz cruzada", wb.sheetnames)

    def test_defect_reports(self):
        incident = self.call("post", "/incidents", 201, json={
            "title": "Total mal calculado", "severity": "Crítica", "steps": "1. Crear factura", "actual": "Total 0"})
        self.assertEqual(self.call("get", "/defects/report?format=pdf").data[:4], b"%PDF")
        wb = load_workbook(io.BytesIO(self.call("get", "/defects/report?format=excel&severity=Crítica").data))
        rows = list(wb.active.iter_rows(min_row=4, values_only=True))
        self.assertTrue(rows and all(r[2] == "Crítica" for r in rows))
        self.assertIn("pérdida de datos", rows[0][3])  # impacto según severidad
        self.assertEqual(self.call("get", f"/incidents/{incident['id']}/report?format=pdf").data[:4], b"%PDF")
        self.call("get", f"/incidents/{incident['id']}/report?format=docx", 400)

    def test_dashboard_metrics_and_menus(self):
        self.assertIn("coverage", self.call("get", "/dashboard/metrics")["docs"])
        html = self.client.get("/dashboard").get_data(as_text=True)
        for view in ("plan", "casos", "trazabilidad", "defectos"):
            self.assertIn(f'id="view-{view}"', html)


if __name__ == "__main__":
    unittest.main(verbosity=2)
